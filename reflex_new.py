﻿
# -*- coding: utf-8 -*-
import os
import pandas as pd
import xlsxwriter
from openpyxl import load_workbook
import re
from scipy.optimize import fsolve,root,brentq
import numpy as np
import matplotlib.pyplot as plt
from scipy.interpolate import BPoly, CubicSpline
import warnings
import xlrd


total_xls_path = './excel/total.xlsx'
file_alpha='./excel/alpha.xlsx'
files_dir = './data/'
lst_files = os.listdir(files_dir)

if not os.path.exists('./excel'):
    os.makedirs('./excel')
if not os.path.exists('./data'):
    os.makedirs('./data/')

def write_xls(file_path):
    # headers for df
    headers1 = ['f', 'alpha_75', 'Y1_75', 'R1_75']
    headers2 = ['alpha_80', 'Y1_80', 'R1_80']
    headers3 = ['alpha_85', 'Y1_85', 'R1_85']
    headers4 = ['alpha_90', 'Y1_90', 'R1_90']
    headers5 = ['alpha_95', 'Y1_95', 'R1_95']
    headers = headers1+headers2+headers3+headers4+headers5
    writer = pd.ExcelWriter(file_path)# https://github.com/PyCQA/pylint/issues/3060 pylint: disable=abstract-class-instantiated
    workbook = writer.book

    for file in lst_files:
        file_name = files_dir+file

        sheet_name = file[:-4]
    # ограниченя экселя((
        data1 = pd.read_csv(file_name, encoding='windows-1251', skiprows=36,
                            nrows=9, header=None, decimal=',', sep=";", usecols=[0, 5, 6, 7],names=headers1)
        data2 = pd.read_csv(file_name, encoding='windows-1251', skiprows=51,
                            nrows=9, header=None, decimal=',', sep=";", usecols=[5, 6, 7],names=headers2)
        data3 = pd.read_csv(file_name, encoding='windows-1251', skiprows=66,
                            nrows=9, header=None, decimal=',', sep=";", usecols=[5, 6, 7],names=headers3)
        data4 = pd.read_csv(file_name, encoding='windows-1251', skiprows=81,
                            nrows=9, header=None, decimal=',', sep=";", usecols=[5, 6, 7],names=headers4)
        data5 = pd.read_csv(file_name, encoding='windows-1251', skiprows=96,
                            nrows=9, header=None, decimal=',', sep=";", usecols=[5, 6, 7],names=headers5)
        data = pd.concat([data1, data2, data3, data4, data5], axis=1)
        data.index=data['f']
        data.to_excel(writer, sheet_name=sheet_name,
                      header=headers, index=False)
        worksheet = writer.sheets[sheet_name]
        # make charts
        chart1 = workbook.add_chart({'type': 'line'})
        chart2 = workbook.add_chart({'type': 'line'})
        chart3 = workbook.add_chart({'type': 'line'})

        for i in range(2, 15, 3):
            chart1.add_series({
                'name':       [sheet_name, 0, i],
                'categories': [sheet_name, 1, 0, 9, 0],
                'values':     [sheet_name, 1, i, 9, i],
                'line':   {'width': 1.25}
            })
            chart2.add_series({
                'name':       [sheet_name, 0, i-1],
                'categories': [sheet_name, 1, 0, 9, 0],
                'values':     [sheet_name, 1, i-1, 9, i-1],
            })
            chart3.add_series({
                'name':       [sheet_name, 0, i+1],
                'categories': [sheet_name, 1, 0, 9, 0],
                'values':     [sheet_name, 1, i+1, 9, i+1],
            })

        chart1.set_title({'name': 'Y','name_font': {'size': 14, 'bold': False}})
        chart1.set_x_axis({
            'name': 'f, Гц',
            'name_font': {'size': 14, 'bold': False},
             })
        chart1.set_style(2)

        chart2.set_title({'name': 'alpha','name_font': {'size': 14, 'bold': False}})
        chart2.set_x_axis({
            'name': 'f, Гц',
            'name_font': {'size': 14, 'bold': False},
             })
        chart2.set_style(2)
        chart3.set_title({'name': 'R','name_font': {'size': 14, 'bold': False}})
        chart3.set_x_axis({
            'name': 'f, Гц',
            'name_font': {'size': 14, 'bold': False},
             })
        chart3.set_style(2)
        # Insert the chart into the worksheet
        worksheet.insert_chart('A19', chart1, {'x_offset': 0, 'y_offset': 0})
        chart1.set_size({'width': 448, 'height': 300})#320x 226
        worksheet.insert_chart('A34', chart2, {'x_offset': 0, 'y_offset': 0})
        chart2.set_size({'width': 448, 'height': 300})#320x 226
        worksheet.insert_chart('A49', chart3, {'x_offset': 0, 'y_offset': 0})
        chart3.set_size({'width': 448, 'height': 300})#320x 226
    # сортировка листов
    workbook.worksheets_objs.sort(key=lambda x: x.name)
    writer.save()
    writer.close()
def find_resonance(file_path,file_alpha):
    xl_file=pd.ExcelFile(file_path)
    wb=xlsxwriter.Workbook(file_alpha)
    sheets=xl_file.sheet_names
    arr_sheets=list(dict.fromkeys([sheets[i] for i in range(len(sheets))]))
    ws1=wb.add_worksheet('Total_alpha_max')
    ws2=wb.add_worksheet('Total_R_min')
    ws3=wb.add_worksheet('Total_Y0')
    total_index_row=1
    for arr_sheet in arr_sheets:
        ws1.write(total_index_row,0,arr_sheet)#запись в листы тотал названий фио
        ws2.write(total_index_row,0,arr_sheet)
        ws3.write(total_index_row,0,arr_sheet)
        merge_format = wb.add_format({'align': 'center'})
        df=pd.read_excel(file_path,sheet_name=arr_sheet,index_col=0)
        df_y=df.iloc[:,[1,4,7,10,13]] #срезы по y,alpha,r
        df_alpha=df.iloc[:,[0,3,6,9,12]]
        df_r=df.iloc[:,[2,5,8,11,14]]
        alpha_names=df_alpha.keys()
        
        max_alpha_values=[]
        max_alpha_frequencis=[]
        roots_y=[]
        min_r_values=[]
        min_r_frequencis=[]
        i=1
        
        print(total_index_row)

        for a_index in range(len(alpha_names)):
            
            x_coords=np.array(df_alpha[df_alpha.columns[a_index]].keys().tolist())
            y_coords_alpha=np.array(df_alpha[df_alpha.columns[a_index]].tolist())
            y_coords_y=np.array(df_y[df_y.columns[a_index]].tolist())
            y_coords_r=np.array(df_r[df_r.columns[a_index]].tolist())
            cubic_alpha=CubicSpline(x_coords,y_coords_alpha)
                #cubic_y=CubicSpline(x_coords,y_coords_y)
            cubic_r=CubicSpline(x_coords,y_coords_r)

            xnew = np.arange(330, 570, 0.1)
            ynew_alpha=cubic_alpha(xnew)
                #ynew_y=cubic_y(xnew)
            ynew_r=cubic_r(xnew)
            bpoly = BPoly.from_derivatives(x_coords,y_coords_y[:,np.newaxis],extrapolate=None)

            max_index_alpha = np.argmax(ynew_alpha)
            max_value_alpha = ynew_alpha[max_index_alpha]
            max_x_alpha = xnew[max_index_alpha]
                
            min_index_r = np.argmin(ynew_r)
            min_value_r = ynew_r[min_index_r]
            min_x_r = xnew[min_index_r]
            # print(y_coords_alpha)
            try:
                root_y = brentq(bpoly, 350, max_x_alpha+50)# Нормуль!
            except ValueError:
                root_y='-'
                # max_x_alpha='-'
                # min_x_r='-'
            max_alpha_values.append('%0.2f' %max_value_alpha)
            max_alpha_frequencis.append('%0.1f' % max_x_alpha)
            roots_y.append(root_y)
            min_r_values.append('%0.2f' %min_value_r)
            min_r_frequencis.append('%0.1f' %min_x_r)

            ws1.write(total_index_row,i,max_x_alpha)
            ws1.write(0,i,alpha_names[i-1][6:]+'dB')
            ws2.write(total_index_row,i,min_x_r)
            ws2.write(0,i,alpha_names[i-1][6:]+'dB')
            ws3.write(total_index_row,i,root_y)
            ws3.write(0,i,alpha_names[i-1][6:]+'dB')
            i+=1
            max_alpha_values.append('%0.2f' %max_value_alpha)
        try:
            for f_y in roots_y:
                if float(f_y)>=float(roots_y[0])+5:
                    report='Частота повысилась'
                    symbol=u'\u2191'
                    break
                elif float(f_y)+5<=float(roots_y[0]):
                    report='Частота понизилась'
                    symbol=u'\u2193'
                    break
                else:
                    report='Частота не изменилась'
                    symbol=u'\u2192'
                    continue
                return f_y,report,symbol
        except ValueError:
            report='-'
            symbol='-'
        ws3.write(total_index_row,7,symbol)
        ws3.write(total_index_row,6,alpha_names[roots_y.index(f_y)][6:])
        for f_alpha in max_alpha_frequencis:
                if float(f_alpha)>float(max_alpha_frequencis[0])+5:
                    report='Частота повысилась'
                    symbol=u'\u2191'
                    break
                elif float(f_alpha)+5<float(max_alpha_frequencis[0]):
                    report='Частота понизилась'
                    symbol=u'\u2193'
                    break
                else:
                    report='Частота не изменилась'
                    symbol=u'\u2192'
                    continue
                return f_alpha,report,symbol
        ws1.write(total_index_row,7,symbol)
        ws1.write(total_index_row,6,alpha_names[max_alpha_frequencis.index(f_alpha)][6:])
        for f_r in min_r_frequencis:
                if float(f_r)>float(min_r_frequencis[0])+5:
                    report='Частота повысилась'
                    symbol=u'\u2191'
                    break
                elif float(f_r)+5<float(min_r_frequencis[0]):
                    report='Частота понизилась'
                    symbol=u'\u2193'
                    break
                else:
                    report='Частота не изменилась'
                    symbol=u'\u2192'
                    continue
                return f_r,report,symbol
        ws2.write(total_index_row,7,symbol)
        ws2.write(total_index_row,6,alpha_names[min_r_frequencis.index(f_r)][6:])
        total_index_row+=1

        print('------')



    wb.close()
write_xls(total_xls_path)
find_resonance(total_xls_path,file_alpha)
print('Done !')
