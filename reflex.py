# -*- coding: utf-8 -*-
import os
import pandas as pd
import xlsxwriter
from openpyxl import load_workbook
import re
from scipy.optimize import fsolve, root, brentq
import numpy as np
import matplotlib.pyplot as plt
from scipy.interpolate import BPoly, CubicSpline
import warnings
import xlrd
import itertools
from collections import Counter

HEADERS = ['f', 'alpha_1', 'Y1_1', 'R1_1', 'alpha_2', 'Y1_2', 'R1_2', 'alpha_3', 'Y1_3', 'R1_3', 'alpha_4', 'Y1_4',
           'R1_4', 'alpha_5', 'Y1_5', 'R1_5']  # TODO: сделать красивше


class Data:
    # Кодировка, путь к файлу, пропуск строк в начале, количество строк в каждом куске df,
    # номер столбцов, пропущенные строки между кусками df, кол-во кусков df
    def __init__(self, encoding, path, skip_rows, num_rows, cols, df_space, df_nums):
        self.encoding = encoding
        self.path = path
        self.skip_rows = skip_rows
        self.num_rows = num_rows
        self.cols = cols
        self.df_space = df_space
        self.df_nums = df_nums

    def create_df(self):
        # TODO: разобраться с ебаными индексами
        result_df = pd.read_csv(self.path, encoding=self.encoding,
                                skiprows=self.skip_rows,
                                nrows=self.num_rows, header=None, decimal=',', sep=";", usecols=[0])

        for i in range(self.df_nums):
            data = pd.read_csv(self.path, encoding=self.encoding,
                               skiprows=self.skip_rows + (self.num_rows + self.df_space) * i,
                               nrows=self.num_rows, header=None, decimal=',', sep=";", usecols=self.cols)
            result_df = pd.concat([result_df, data], axis=1, ignore_index=False)
        result_df.columns = HEADERS
        result_df.index = result_df['f']

        return result_df


# Функция возвращает словарь сокращенное имя файла - путь к файлу с проходом по папкам
def receive_files_path(path):
    files_list = []
    files_names = []

    for root, dirs, files in os.walk(path):
        for f in files:
            file_path = os.path.abspath(os.path.join(root, f))
            file_name = file_path.split('/')[-1][:-4]
            if f.endswith('.csv'):
                files_list.append(file_path)
                files_names.append(file_name)
    files_dict = dict(zip(files_names, files_list))
    return files_dict


# Функция записи data frame, charts в xls
def write_data_xls(result_df, sheet_name):
    # TODO: возможно разбить
    chart_name = {'y': {'name': 'Y', 'name_font': {'size': 14, 'bold': False}},
                  'alpha': {'name': 'alpha', 'name_font': {'size': 14, 'bold': False}},
                  'r': {'name': 'R', 'name_font': {'size': 14, 'bold': False}}}
    chart_axis = {
        'name': 'f, Гц',
        'name_font': {'size': 14, 'bold': False},
    }

    result_df.to_excel(writer, sheet_name=sheet_name,  # TODO: rename sheet_name
                       index=False)
    draw_chart(writer, sheet_name, 2, 15, 3, 1, 11, chart_name['y'], chart_axis, 'A19')
    draw_chart(writer, sheet_name, 1, 14, 3, 1, 11, chart_name['alpha'], chart_axis, 'A34')
    draw_chart(writer, sheet_name, 3, 16, 3, 1, 11, chart_name['r'], chart_axis, 'A49')


def write_total_xls(data_folder_path, result_df, name):
    receive_files_path(data_folder_path)
    resonance_list = find_resonance(result_df)
    changes_list = find_changes(result_df)
    print(changes_list)
    ws1.write(row_index, 0, name)
    ws2.write(row_index, 0, name)
    ws3.write(row_index, 0, name)
    column_index = 1

    for i in range(len(resonance_list[0])):
        ws1.write(0, column_index, i + 1)
        ws2.write(0, column_index, i + 1)
        ws3.write(0, column_index, i + 1)
        ws1.write(row_index, column_index, resonance_list[0][i])
        ws2.write(row_index, column_index, resonance_list[1][i])
        ws3.write(row_index, column_index, resonance_list[2][i])
        column_index += 1


def draw_chart(writer, sheet_name, start_col, end_col, step_col, start_row, end_row, chart_name, chart_axis,
               chart_point):
    workbook = writer.book
    chart = workbook.add_chart({'type': 'line'})
    worksheet = writer.sheets[sheet_name]
    for i in range(start_col, end_col, step_col):  # 2,15,3
        chart.add_series({
            'name': [sheet_name, 0, i],
            'categories': [sheet_name, start_row, 0, end_row, 0],  # 1,11
            'values': [sheet_name, start_row, i, end_row, i],
            'line': {'width': 1.25}
        })
    chart.set_title(chart_name)
    chart.set_x_axis(chart_axis)
    chart.set_style(2)
    worksheet.insert_chart(chart_point, chart, {'x_offset': 0, 'y_offset': 0})
    chart.set_size({'width': 448, 'height': 300})  # 320x 226


def find_changes(result_df):
    # наибольшее количество вхождений
    resonance_list = find_resonance(result_df)
    resonance_alpha = list(map(float, resonance_list[0]))
    resonance_alpha_arr = np.array(resonance_alpha)
    z_alpha = resonance_alpha_arr[:, None] >= resonance_alpha_arr + 3  # 3 db matrix
    alpha = np.where(z_alpha == True)

    resonance_r = list(map(float, resonance_list[1]))
    resonance_r_arr = np.array(resonance_r)
    z_r = resonance_r_arr[:, None] >= resonance_r_arr + 3  # 3 db matrix
    r = np.where(z_r == True)

    resonance_y= list(map(float, resonance_list[2]))
    resonance_y_arr = np.array(resonance_y)
    z_y = resonance_y_arr[:, None] >= resonance_y_arr + 3  # 3 db matrix
    y = np.where(z_y == True)
    # t = list(zip(a[0], a[1]))
    most_common_alpha = Counter(alpha[0]).most_common(1)[0][0]
    most_common_r = Counter(r[0]).most_common(1)[0][0]
    most_common_y = Counter(y[0]).most_common(1)[0][0]
    #первое вхождение
    print(resonance_r)
    first_change_alpha = np.where(z_alpha == True)[0][0]
    first_change_r = np.where(z_r == True)[0][0]
    first_change_y = np.where(z_y == True)[0][0]
    print(first_change_r)

    return first_change_alpha,first_change_r,first_change_y,\
           most_common_alpha,most_common_r,most_common_y


def find_resonance(result_df):
    df_alpha = result_df.iloc[:, [1, 4, 7, 10, 13]]  # срезы по y,alpha,r
    df_y = result_df.iloc[:, [2, 5, 8, 11, 14]]
    df_r = result_df.iloc[:, [3, 6, 9, 12, 15]]
    num_columns = len(df_alpha.columns)
    frequencies = list(result_df.index.values)
    frequencies_cubic = np.arange(frequencies[0], frequencies[-1], 0.1)
    max_alpha_values = []
    max_alpha_frequencis = []
    roots_y = []
    min_r_values = []
    min_r_frequencis = []
    for index in range(num_columns):
        alpha_values = np.array(df_alpha[df_alpha.columns[index]].tolist())
        y_values = np.array(df_y[df_y.columns[index]].tolist())
        r_values = np.array(df_r[df_r.columns[index]].tolist())
        # cubic spline functions
        cubic_alpha = CubicSpline(frequencies, alpha_values)
        cubic_r = CubicSpline(frequencies, r_values)
        # new y axis values for alpha,r
        alpha_values_cubic = cubic_alpha(frequencies_cubic)
        r_values_cubic = cubic_r(frequencies_cubic)
        # maximum alpha
        max_alpha_index = np.argmax(alpha_values_cubic)
        max_alpha_value = alpha_values_cubic[max_alpha_index]
        max_alpha_frequency = frequencies_cubic[max_alpha_index]
        bpoly = BPoly.from_derivatives(frequencies, y_values[:, np.newaxis], extrapolate=None)
        # minimum R values
        min_r_index = np.argmin(r_values_cubic)
        min_r_value = r_values_cubic[min_r_index]
        min_r_frequency = frequencies_cubic[min_r_index]
        # приближение в 50 Гц для поиска перечения Y
        frequency_approximation = 50
        try:
            root_y = brentq(bpoly, frequencies[0], max_alpha_frequency + frequency_approximation)  # Нормуль!
        except ValueError:
            root_y = 'NaN'
        max_alpha_values.append('%0.2f' % max_alpha_value)
        max_alpha_frequencis.append('%0.1f' % max_alpha_frequency)
        roots_y.append(root_y)
        min_r_values.append('%0.2f' % min_r_value)
        min_r_frequencis.append('%0.1f' % min_r_frequency)
    return max_alpha_frequencis, min_r_frequencis, roots_y


writer = pd.ExcelWriter('excel/graph.xlsx')
files_dict = receive_files_path('data')
wb = xlsxwriter.Workbook('excel/total.xlsx')
ws1 = wb.add_worksheet('Total_Y')
ws2 = wb.add_worksheet('Total_alpha_max')
ws3 = wb.add_worksheet('Total_R_min')

row_index = 1
for key, value in files_dict.items():
    data = Data('windows-1251', value, 35, 11, [5, 6, 7], 4, 5)
    df = data.create_df()
    write_data_xls(df, key)
    # find_changes(df)
    write_total_xls('data', df, key)
    row_index += 1

writer.save()
wb.close()
